#!/usr/bin/env python3
"""
build_deck.py — Build a presentation on the LF 2025 Template (Google Slides
export) from a JSON spec, optionally re-skinned with a foundation's brand kit.

    python3 build_deck.py spec.json output.pptx [--brand brand.json] [--max-slides N]

The script opens assets/LF-2025-Template.pptx, drops the sample slides, and
adds one slide per spec entry using ONLY the template's own layouts, so fonts,
sizes, gradient backgrounds, footer bar, logo and slide numbers are inherited.
With --brand it then swaps the LF logo for the project logo, maps the LF
palette to the project palette and the template fonts to the project fonts
(see references/brand-override.md).

Spec:
{
  "slides": [
    {"layout": "title", "title": "…", "subtitle": "…", "notes": "…"},
    {"layout": "section", "title": "…"},
    {"layout": "statement", "title": "Big claim", "body": ["support", "support"]},
    {"layout": "content", "title": "…", "bullets": ["**Lead —** rest", ["sub-bullet", 1]]},
    {"layout": "two_col", "title": "…", "left": ["…"], "right": ["…"]},
    {"layout": "content_image", "title": "…", "bullets": ["…"], "image": "chart.png"},
    {"layout": "logos", "title": "Members", "images": ["a.png", "b.png", "c.png", "d.png"]},
    {"layout": "table", "title": "…", "table": {"header": ["A", "B"], "rows": [["1", "2"]]}},
    {"layout": "chart", "title": "…", "image": "chart.png", "source": "Source: LFX Insights, Sep 2026"},
    {"layout": "closing"}
  ]
}

Bullets: a string is a level-0 bullet; ["text", level] indents (0–2).
**bold** inside any text becomes a bold run. Any slide may carry "notes".

Layout keys → template layout:
  title            TITLE                    dark title slide (title + subtitle)
  title_light      TITLE_2                  light title slide
  title_image      TITLE_1                  dark title + square image (right)
  section          SECTION_HEADER           dark section divider
  section_light    SECTION_HEADER_2         light section divider
  statement        SECTION_HEADER_1         dark: big statement left, body right
  statement_image  TITLE_1_1                dark: statement left, small square image right
  body_image       TITLE_1_1_1              dark: body left, tall image right
  logos            SECTION_HEADER_3         dark: title + 4 square logos
  logos_light      SECTION_HEADER_3_1       light: title + 4 square logos
  content          TITLE_AND_BODY           light: title + bullets (default)
  content_dense    TITLE_AND_BODY_1         light: denser bullets (appendix)
  two_col          TITLE_AND_TWO_COLUMNS    light: title + two bullet columns
  two_col_dense    TITLE_AND_BODY_1_1       light: dense two columns (appendix)
  content_image    TITLE_AND_TWO_COLUMNS_1  light: bullets left, square image right
  table            TITLE_ONLY               light: title + template-style table
  chart            TITLE_ONLY               light: title + image scaled into body
  title_only       TITLE_ONLY               light: title only
  blank            BLANK
  closing          TITLE_2_1                LF (or project) logo closing slide
"""
import argparse
import json
import re
import shutil
import sys
import tempfile
import zipfile
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

sys.path.insert(0, str(Path(__file__).resolve().parent))
import lfbrand  # noqa: E402

HERE = Path(__file__).resolve().parent
TEMPLATE = HERE.parent / "assets" / "LF-2025-Template.pptx"

LAYOUTS = {
    "title": "TITLE", "title_light": "TITLE_2", "title_image": "TITLE_1",
    "section": "SECTION_HEADER", "section_light": "SECTION_HEADER_2",
    "statement": "SECTION_HEADER_1", "statement_image": "TITLE_1_1",
    "body_image": "TITLE_1_1_1", "logos": "SECTION_HEADER_3",
    "logos_light": "SECTION_HEADER_3_1", "content": "TITLE_AND_BODY",
    "content_dense": "TITLE_AND_BODY_1", "two_col": "TITLE_AND_TWO_COLUMNS",
    "two_col_dense": "TITLE_AND_BODY_1_1", "content_image": "TITLE_AND_TWO_COLUMNS_1",
    "table": "TITLE_ONLY", "chart": "TITLE_ONLY", "title_only": "TITLE_ONLY",
    "blank": "BLANK", "closing": "TITLE_2_1",
}

# Body area on TITLE_ONLY (matches TITLE_AND_BODY's body placeholder)
BODY_BOX = (Inches(0.34), Inches(1.26), Inches(9.32), Inches(3.74))
TABLE_HEADER_FILL = "4285F4"   # template agenda-table header (slide 13)
TABLE_ZEBRA_FILL = "EEEEEE"

BOLD_RE = re.compile(r"\*\*(.+?)\*\*")


def remove_all_slides(prs):
    sldIdLst = prs.slides._sldIdLst
    for sldId in list(sldIdLst):
        prs.part.drop_rel(sldId.rId)
        sldIdLst.remove(sldId)


def layout_by_name(prs, name):
    for layout in prs.slide_layouts:
        if layout.name == name:
            return layout
    raise SystemExit(f"Template has no layout named {name!r}")


def placeholders(slide):
    out = {"title": [], "body": [], "pic": [], "other": []}
    for ph in slide.placeholders:
        tname = str(ph.placeholder_format.type).split(".")[-1].split(" ")[0]
        if tname in ("TITLE", "CENTER_TITLE", "SUBTITLE"):
            out["title"].append(ph)
        elif tname in ("BODY", "OBJECT"):
            out["body"].append(ph)
        elif tname == "PICTURE":
            out["pic"].append(ph)
        elif tname == "SLIDE_NUMBER":
            continue
        else:
            out["other"].append(ph)
    for k in out:
        out[k].sort(key=lambda p: p.placeholder_format.idx)
    return out


def set_runs(paragraph, text):
    """Write text into a paragraph, turning **bold** spans into bold runs."""
    text = str(text)
    pos = 0
    first = True
    for m in BOLD_RE.finditer(text):
        if m.start() > pos:
            r = paragraph.add_run(); r.text = text[pos:m.start()]
        r = paragraph.add_run(); r.text = m.group(1); r.font.bold = True
        pos = m.end(); first = False
    if pos < len(text) or first:
        r = paragraph.add_run(); r.text = text[pos:]


def fill_text(ph, text):
    tf = ph.text_frame
    tf.clear()
    set_runs(tf.paragraphs[0], text)


def fill_bullets(ph, bullets):
    tf = ph.text_frame
    tf.clear()
    first = True
    for b in bullets:
        text, level = (b[0], int(b[1])) if isinstance(b, (list, tuple)) else (b, 0)
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        set_runs(p, text)
        p.level = max(0, min(level, 2))


def insert_picture(ph, path):
    if not Path(path).exists():
        print(f"  ! image not found, leaving placeholder empty: {path}", file=sys.stderr)
        return
    try:
        ph.insert_picture(str(path))
    except Exception as e:
        print(f"  ! could not insert {path}: {e}", file=sys.stderr)


def add_table(slide, spec, brand=None):
    header = spec.get("header", [])
    rows = spec.get("rows", [])
    ncols = max(len(header), max((len(r) for r in rows), default=0))
    nrows = len(rows) + (1 if header else 0)
    left, top, width, height = BODY_BOX
    height = min(height, Inches(0.42) * nrows)
    shape = slide.shapes.add_table(nrows, ncols, left, top, width, height)
    tbl = shape.table
    # remove the default PowerPoint table style flags so our fills show cleanly
    tblPr = tbl._tbl.tblPr
    for attr in ("firstRow", "bandRow"):
        if tblPr.get(attr) is not None:
            tblPr.set(attr, "0")
    hdr_fill = TABLE_HEADER_FILL
    if brand:
        hdr_fill = lfbrand.table_header_fill(brand)
    r0 = 0
    if header:
        for c, val in enumerate(header):
            cell = tbl.cell(0, c)
            cell.text = str(val)
            cell.fill.solid(); cell.fill.fore_color.rgb = RGBColor.from_string(hdr_fill)
            for p in cell.text_frame.paragraphs:
                for run in p.runs:
                    run.font.bold = True; run.font.size = Pt(13)
                    run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        r0 = 1
    for r, row in enumerate(rows):
        fill = TABLE_ZEBRA_FILL if r % 2 == 1 else "FFFFFF"
        for c in range(ncols):
            val = row[c] if c < len(row) else ""
            cell = tbl.cell(r0 + r, c)
            cell.text = str(val)
            cell.fill.solid(); cell.fill.fore_color.rgb = RGBColor.from_string(fill)
            for p in cell.text_frame.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(12)
                    run.font.color.rgb = RGBColor(0x22, 0x22, 0x22)


def add_body_image(slide, path):
    if not Path(path).exists():
        print(f"  ! image not found: {path}", file=sys.stderr)
        return
    left, top, width, height = BODY_BOX
    pic = slide.shapes.add_picture(str(path), left, top, height=height)
    if pic.width < width:
        pic.left = left + (width - pic.width) // 2
    elif pic.width > width:
        ratio = width / pic.width
        pic.width = width
        pic.height = int(pic.height * ratio)


def add_source_line(slide, text):
    left, top, width, height = BODY_BOX
    tb = slide.shapes.add_textbox(left, top + height + Inches(0.02), width, Inches(0.25))
    p = tb.text_frame.paragraphs[0]
    r = p.add_run(); r.text = text
    r.font.size = Pt(8); r.font.italic = True
    r.font.color.rgb = RGBColor(0xA1, 0xA1, 0xA4)
    from pptx.enum.text import PP_ALIGN
    p.alignment = PP_ALIGN.RIGHT


def strip_empty_placeholders(slide):
    for ph in list(slide.placeholders):
        tname = str(ph.placeholder_format.type).split(".")[-1].split(" ")[0]
        if tname == "SLIDE_NUMBER":
            continue
        if tname == "PICTURE":
            if ph.element.tag.endswith("}pic"):
                continue
            ph.element.getparent().remove(ph.element)
            continue
        if ph.has_text_frame and not ph.text_frame.text.strip():
            ph.element.getparent().remove(ph.element)


# ---------------------------------------------------------------- brand skin

def apply_brand(pptx_path, brand):
    """Post-process the saved .pptx: swap logo media, recolor, refont."""
    color_map = lfbrand.color_map_for_slides(brand)
    font_map = lfbrand.font_map_for_slides(brand)
    tmpdir = Path(tempfile.mkdtemp())
    logo_media = {}
    # template media: image1.png = LF color logo (light bg), image2.png = white logo (dark bg)
    if brand["logo"].get("light"):
        logo_media["ppt/media/image1.png"] = lfbrand.fit_logo(brand["logo"]["light"], 862, 284, tmpdir / "l.png")
    if brand["logo"].get("dark"):
        logo_media["ppt/media/image2.png"] = lfbrand.fit_logo(brand["logo"]["dark"], 862, 284, tmpdir / "d.png")
    elif brand["logo"].get("light"):
        # no dark variant supplied: reuse the light logo on dark slides (better than the LF mark there)
        logo_media["ppt/media/image2.png"] = logo_media["ppt/media/image1.png"]

    src = Path(pptx_path)
    out = src.with_suffix(".branded.tmp")
    with zipfile.ZipFile(src) as zin, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename in logo_media:
                data = Path(logo_media[item.filename]).read_bytes()
            elif item.filename.endswith(".xml") and (
                item.filename.startswith("ppt/slideMasters/") or
                item.filename.startswith("ppt/slideLayouts/") or
                item.filename.startswith("ppt/slides/") or
                item.filename.startswith("ppt/theme/")
            ):
                xml = data.decode("utf-8")
                xml = lfbrand.replace_colors_in_xml(xml, color_map)
                xml = lfbrand.replace_fonts_in_xml(xml, font_map)
                data = xml.encode("utf-8")
            zout.writestr(item, data)
    shutil.move(str(out), str(src))
    shutil.rmtree(tmpdir, ignore_errors=True)


# --------------------------------------------------------------------- build

def build(spec_path, out_path, brand_path=None, max_slides=None):
    spec = json.loads(Path(spec_path).read_text())
    slides = spec["slides"]
    if max_slides and len(slides) > max_slides:
        raise SystemExit(f"Deck has {len(slides)} slides; the limit is {max_slides}.")
    brand = lfbrand.load_brand(brand_path or spec.get("brand"))

    prs = Presentation(str(TEMPLATE))
    remove_all_slides(prs)

    for i, s in enumerate(slides, 1):
        key = s.get("layout", "content")
        if key not in LAYOUTS:
            raise SystemExit(f"Slide {i}: unknown layout key {key!r}. Valid: {sorted(LAYOUTS)}")
        slide = prs.slides.add_slide(layout_by_name(prs, LAYOUTS[key]))
        ph = placeholders(slide)

        if s.get("title") and ph["title"]:
            fill_text(ph["title"][0], s["title"])
        if s.get("subtitle") and len(ph["title"]) > 1:
            fill_text(ph["title"][1], s["subtitle"])

        if key in ("two_col", "two_col_dense"):
            if ph["body"]:
                fill_bullets(ph["body"][0], s.get("left", []))
            if len(ph["body"]) > 1:
                fill_bullets(ph["body"][1], s.get("right", []))
        else:
            body = s.get("bullets") or s.get("body")
            if body and ph["body"]:
                fill_bullets(ph["body"][0], body)

        imgs = s.get("images") or ([s["image"]] if s.get("image") and key != "chart" else [])
        for pic_ph, img in zip(ph["pic"], imgs):
            insert_picture(pic_ph, img)

        if key == "table" and s.get("table"):
            add_table(slide, s["table"], brand)
        if key == "chart" and s.get("image"):
            add_body_image(slide, s["image"])
        if s.get("source"):
            add_source_line(slide, s["source"])
        if s.get("notes"):
            slide.notes_slide.notes_text_frame.text = s["notes"]

        strip_empty_placeholders(slide)

    prs.save(out_path)
    if brand:
        apply_brand(out_path, brand)
    print(f"Wrote {out_path} — {len(slides)} slides on LF 2025 Template; brand: {lfbrand.summarize(brand)}")


if __name__ == "__main__":
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("spec"); ap.add_argument("output")
    ap.add_argument("--brand", help="brand.json for a foundation-specific skin")
    ap.add_argument("--max-slides", type=int, default=None)
    a = ap.parse_args()
    build(a.spec, a.output, a.brand, a.max_slides)
