#!/usr/bin/env python3
"""
build_doc.py — Build a Word document on the LF Agent DOCS Template (or a
letter on the LF Letterhead) from one Markdown file with front matter,
optionally re-skinned with a foundation's brand kit.

    python3 build_doc.py input.md output.docx [--brand brand.json] [--letterhead]
                         [--no-cover] [--strip-fonts] [--body-size 9]

The rendering engine is build_lf_doc.py (the renderer the foundation agents
already use: cover page, header/footer tokens, Roboto Slab headings, template
tables with brand-colored header rows, bullets, decimal numbering, code
blocks, blockquotes, rules, <<<PAGEBREAK>>>). This wrapper adds:

  * front matter instead of a long CLI (all build_lf_doc options are keys);
  * brand.json (see references/brand-override.md) → primary / secondary
    colors + cover logo, AND the template re-skin build_lf_doc does not do:
    the header logo is swapped for the project logo and the template fonts
    are swapped for the brand fonts;
  * --letterhead: business-letter layout on the LF Letterhead template.

FRONT MATTER (document):
    ---
    project_name: PyTorch Foundation          # header + cover title      (required unless brand.json has it)
    document_type: Message Foundation          # header + cover            (required)
    document_status: DRAFT for ED review       # header + cover
    other_details: v0.3                        # header
    platform: LFX Marketing OS                 # footer (default)
    agent_source: Message Foundation Agent     # footer
    date: September 22, 2026                   # footer + cover (default today)
    subtitle: One of three LFX Marketing OS foundational documents   # cover
    details:                                   # cover details table rows (Label=Value)
      - Repository=https://github.com/pytorch
      - Prepared for=PyTorch Foundation leadership
    logo: pytorch-logo.png                     # cover logo when no brand.json (brand.json logo.light wins)
    primary: "#EE4C2C"                         # colors when no brand.json
    secondary: "#2B2B2B"
    cover: true                                # false = --no-cover
    brand: ./brand.json                        # same as --brand
    ---
    # 1. First section …

Body Markdown: #…##### headings, paragraphs (**bold**, *italic*, `code`,
[links](url)), - bullets and 1. numbered lists (two-space nesting), GFM pipe
tables (<br> for in-cell breaks), > blockquotes, fenced code blocks, --- rule,
a line of exactly <<<PAGEBREAK>>> (or ---pagebreak--- / \\pagebreak) for a
page break. Do not put the title or header/footer text in the body.

FRONT MATTER (letter, --letterhead): date, to (multi-line with "|"),
salutation, closing, signature (multi-line). Body = the letter paragraphs.
"""
import argparse
import copy
import datetime as dt
import re
import shutil
import sys
import tempfile
import zipfile
from pathlib import Path
from types import SimpleNamespace

import docx
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

sys.path.insert(0, str(Path(__file__).resolve().parent))
import lfbrand  # noqa: E402
import build_lf_doc as engine  # noqa: E402

HERE = Path(__file__).resolve().parent
ASSETS = HERE.parent / "assets"
TEMPLATE_DOCS = ASSETS / "LF-Agent-DOCS-Template.docx"
TEMPLATE_LETTER = ASSETS / "LF-Letterhead-Template.docx"
LETTER_PT = 10


# ------------------------------------------------------------ front matter

def parse_front_matter(text):
    meta = {}
    if text.startswith("---"):
        end = text.find("\n---", 3)
        if end != -1:
            block = text[3:end].strip("\n")
            key = None
            for line in block.splitlines():
                if re.match(r"^\s+-\s+", line) and key:                 # list item
                    meta.setdefault(key, [])
                    if isinstance(meta[key], str):
                        meta[key] = [meta[key]] if meta[key] else []
                    meta[key].append(re.sub(r"^\s+-\s+", "", line).strip().strip('"'))
                    continue
                if re.match(r"^\s+", line) and key:                       # continuation
                    if isinstance(meta[key], list):
                        continue
                    meta[key] = (meta[key] + "\n" + line.strip()).strip()
                    continue
                m = re.match(r"^([A-Za-z_][\w-]*):\s*(.*)$", line)
                if m:
                    key, val = m.group(1), m.group(2).strip().strip('"').strip("'")
                    meta[key] = "" if val == "|" else val
            text = text[end + 4:]
    return meta, text.lstrip("\n")


def normalize_breaks(md):
    return re.sub(r"^\s*(---pagebreak---|\\pagebreak|<!-- pagebreak -->)\s*$", "<<<PAGEBREAK>>>", md, flags=re.M)


def as_bool(v, default=True):
    if v is None or v == "":
        return default
    return str(v).strip().lower() not in ("false", "no", "0", "off")


# -------------------------------------------------------------------- brand

def apply_brand_skin(docx_path, brand):
    """Post-process: swap the header logo, map template fonts to brand fonts,
    recolor the letterhead's Azure footer text to the brand primary."""
    color_map = {}
    if brand["palette"]["primary"] != "0094FF":
        color_map["0094ff"] = brand["palette"]["primary"]
    font_map = lfbrand.font_map_for_docs(brand)
    tmpdir = Path(tempfile.mkdtemp())
    logo = None
    if brand["logo"].get("light"):
        logo = lfbrand.fit_logo(brand["logo"]["light"], 1000, 329, tmpdir / "logo.png")
    if not (logo or color_map or font_map):
        return
    src = Path(docx_path); out = src.with_suffix(".branded.tmp")
    with zipfile.ZipFile(src) as zin, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if logo and item.filename == "word/media/image1.png":          # template header logo
                data = Path(logo).read_bytes()
            elif item.filename.endswith(".xml") and item.filename.startswith("word/") and "fontTable" not in item.filename:
                xml = data.decode("utf-8")
                xml = lfbrand.replace_colors_in_xml(xml, color_map)
                xml = lfbrand.replace_fonts_in_xml(xml, font_map)
                data = xml.encode("utf-8")
            zout.writestr(item, data)
    shutil.move(str(out), str(src)); shutil.rmtree(tmpdir, ignore_errors=True)


# ----------------------------------------------------------------- document

def build_document(meta, body_md, out_path, brand, opts):
    project = meta.get("project_name") or (brand["project_name"] if brand else None)
    if not project:
        raise SystemExit("front matter needs project_name (or brand.json with project_name)")
    if not meta.get("document_type"):
        raise SystemExit("front matter needs document_type")

    primary = brand["palette"]["primary"] if brand else engine.norm_hex(meta.get("primary"))
    secondary = brand["palette"]["secondary"] if brand else engine.norm_hex(meta.get("secondary"))
    if brand and primary == "0094FF" and secondary == "003778" and not (meta.get("primary")):
        primary, secondary = None, None          # LF-default palette → template neutral, as build_lf_doc does
    logo = (brand["logo"].get("light") if brand else None) or meta.get("logo") or None

    details = meta.get("details") or []
    if isinstance(details, str):
        details = [details] if details else []

    args = SimpleNamespace(
        project=project,
        doc_type=meta["document_type"],
        status=meta.get("document_status", "Draft v1 — for review"),
        other=meta.get("other_details", ""),
        subtitle=meta.get("subtitle", ""),
        agent=meta.get("agent_source", ""),
        platform=meta.get("platform", "LFX Marketing OS"),
        date=meta.get("date") or dt.date.today().strftime("%B %d, %Y"),
        detail=details or None,
    )
    body_pt = float(opts.body_size or meta.get("body_size") or engine.BODY_PT)

    document = docx.Document(str(opts.template or TEMPLATE_DOCS))
    engine.clear_body(document)
    engine.fill_header_footer(document, {
        "Project_Name": args.project, "Document_Type": args.doc_type,
        "Document_Status": args.status, "Other_Details": args.other,
        "Platform": args.platform, "Agent_Source": args.agent, "Date": args.date,
    })
    logo_png = engine.prepare_logo(logo) if logo else None
    if as_bool(meta.get("cover"), True) and not opts.no_cover:
        engine.add_cover(document, args, (primary, secondary), body_pt, logo_png)
    engine.Renderer(document, (primary, secondary), body_pt).render(engine.parse_blocks(normalize_breaks(body_md)))
    document.save(out_path)
    if opts.strip_fonts:
        engine.strip_embedded_fonts(out_path)
    if brand:
        apply_brand_skin(out_path, brand)
    print(f"Wrote {out_path} on LF Agent DOCS Template; brand: {lfbrand.summarize(brand)}")


# ------------------------------------------------------------------- letter

def _line(document, text, size_pt, tight=False):
    p = document.add_paragraph()
    if tight:
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.15
    engine.add_inline(p, text, size_pt=size_pt)
    return p


def build_letter(meta, body_md, out_path, brand, opts):
    document = docx.Document(str(opts.template or TEMPLATE_LETTER))
    engine.clear_body(document)
    size_pt = float(opts.body_size or meta.get("body_size") or LETTER_PT)

    if brand and brand.get("footer_address"):
        lines = [ln for ln in brand["footer_address"].splitlines() if ln.strip()]
        paras = [p for p in document.sections[0].footer.paragraphs if p.text.strip()]
        for p, ln in zip(paras, lines):
            for r in p.runs[1:]:
                r.text = ""
            p.runs[0].text = ln
        for p in paras[len(lines):]:
            for r in p.runs:
                r.text = ""

    _line(document, meta.get("date") or dt.date.today().strftime("%B %-d, %Y"), size_pt)
    if meta.get("to"):
        document.add_paragraph()
        for ln in str(meta["to"]).splitlines():
            _line(document, ln, size_pt, tight=True)
    document.add_paragraph()
    if meta.get("salutation"):
        _line(document, meta["salutation"], size_pt)

    primary = brand["palette"]["primary"] if brand and brand["palette"]["primary"] != "0094FF" else None
    r = engine.Renderer(document, (primary, None), size_pt)
    r.usable_width_in = 7.0                     # letterhead margins are 0.75in
    r.render(engine.parse_blocks(normalize_breaks(body_md)))

    if meta.get("closing"):
        document.add_paragraph(); _line(document, meta["closing"], size_pt)
    if meta.get("signature"):
        document.add_paragraph()
        for ln in str(meta["signature"]).splitlines():
            _line(document, ln, size_pt, tight=True)

    document.save(out_path)
    if brand:
        apply_brand_skin(out_path, brand)
    print(f"Wrote {out_path} on LF Letterhead; brand: {lfbrand.summarize(brand)}")


# -------------------------------------------------------------------- main

def build(md_path, out_path, brand_path=None, letterhead=False, **kw):
    opts = SimpleNamespace(body_size=None, no_cover=False, strip_fonts=False, template=None)
    for k, v in kw.items():
        setattr(opts, k, v)
    text = Path(md_path).read_text(encoding="utf-8")
    meta, body_md = parse_front_matter(text)
    brand_file = brand_path or meta.get("brand")
    if brand_file and not Path(brand_file).is_absolute():
        cand = Path(md_path).parent / brand_file
        brand_file = str(cand) if cand.exists() else brand_file
    brand = lfbrand.load_brand(brand_file)
    Path(out_path).parent.mkdir(parents=True, exist_ok=True)
    if letterhead:
        build_letter(meta, body_md, out_path, brand, opts)
    else:
        build_document(meta, body_md, out_path, brand, opts)


if __name__ == "__main__":
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("markdown"); ap.add_argument("output")
    ap.add_argument("--brand", help="brand.json for a foundation-specific skin")
    ap.add_argument("--letterhead", action="store_true", help="business letter on the LF Letterhead template")
    ap.add_argument("--no-cover", action="store_true", help="skip the cover page (documents)")
    ap.add_argument("--strip-fonts", action="store_true", help="drop embedded fonts (small file for Google Drive upload)")
    ap.add_argument("--body-size", type=float, default=None, help="body text size in pt (template: 9 doc / 10 letter)")
    ap.add_argument("--template", default=None, help="override the template .docx")
    a = ap.parse_args()
    build(a.markdown, a.output, a.brand, a.letterhead,
          no_cover=a.no_cover, strip_fonts=a.strip_fonts, body_size=a.body_size, template=a.template)
